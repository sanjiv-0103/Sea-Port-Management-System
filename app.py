from flask import Flask, render_template, request, redirect, session, url_for, flash, send_file
import mysql.connector
from openpyxl import Workbook
from io import BytesIO

app = Flask(__name__)
app.secret_key = "seaport_secret_key"


def get_db_connection():
    return mysql.connector.connect(
        host="localhost",
        user="root",
        password="S@njivKrish2006",          # add MySQL password here if you have one
        database="dbms"                      # change if your DB name is different
    )


ENTITIES = {
    "bookings": {
        "title": "Shipment Bookings",
        "table": "shipment_booking",
        "pk": "booking_id",
        "columns": ["booking_id", "booking_date", "booking_status"]
    },
    "customers": {
        "title": "Customers",
        "table": "customer_3nf",
        "pk": "customer_id",
        "columns": ["customer_id", "customer_name", "customer_contact"]
    },
    "voyages": {
        "title": "Voyages",
        "table": "voyage_3nf",
        "pk": "voyage_id",
        "columns": ["voyage_id", "trip_date", "route_id"]
    },
    "routes": {
        "title": "Sea Routes",
        "table": "sea_route_3nf",
        "pk": "route_id",
        "columns": ["route_id", "route_name"]
    },
    "payments": {
        "title": "Payments",
        "table": "payment_3nf",
        "pk": "payment_id",
        "columns": ["payment_id", "amount", "payment_date"]
    },
    "containers": {
        "title": "Containers",
        "table": "container_2nf",
        "pk": "container_id",
        "columns": ["container_id", "total_weight"]
    },
    "cargo": {
        "title": "Cargo",
        "table": "cargo_2nf",
        "pk": "cargo_id",
        "columns": ["cargo_id", "cargo_description"]
    }
}


def is_logged_in():
    return "user" in session


@app.route("/")
def home():
    if is_logged_in():
        return redirect("/dashboard")
    return redirect("/login")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form["username"]
        password = request.form["password"]

        db = get_db_connection()
        cursor = db.cursor(dictionary=True)

        cursor.execute(
            "SELECT * FROM login_users WHERE username=%s AND password=%s",
            (username, password)
        )
        user = cursor.fetchone()

        cursor.close()
        db.close()

        if user:
            session["user"] = username
            flash("Login successful.")
            return redirect("/dashboard")
        else:
            flash("Invalid username or password.")

    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("Logged out successfully.")
    return redirect("/login")


@app.route("/dashboard")
def dashboard():
    if not is_logged_in():
        return redirect("/login")

    db = get_db_connection()
    cursor = db.cursor()

    counts = {}

    for key, entity in ENTITIES.items():
        cursor.execute(f"SELECT COUNT(*) FROM {entity['table']}")
        counts[key] = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM shipment_booking WHERE booking_status='Confirmed'")
    counts["confirmed"] = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM shipment_booking WHERE booking_status='Pending'")
    counts["pending"] = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM shipment_booking WHERE booking_status='Cancelled'")
    counts["cancelled"] = cursor.fetchone()[0]

    status_labels = []
    status_values = []

    cursor.execute(
        "SELECT booking_status, COUNT(*) FROM shipment_booking GROUP BY booking_status"
    )
    status_data = cursor.fetchall()

    for status, count in status_data:
        status_labels.append(status)
        status_values.append(count)

    cursor.execute("SELECT * FROM activity_log ORDER BY created_at DESC LIMIT 5")
    activities = cursor.fetchall()

    cursor.close()
    db.close()

    return render_template(
        "dashboard.html",
        counts=counts,
        status_labels=status_labels,
        status_values=status_values,
        activities=activities
    )


@app.route("/entity/<entity_key>")
def view_entity(entity_key):
    if not is_logged_in():
        return redirect("/login")

    if entity_key not in ENTITIES:
        flash("Invalid menu option.")
        return redirect("/dashboard")

    entity = ENTITIES[entity_key]

    db = get_db_connection()
    cursor = db.cursor(dictionary=True)

    cursor.execute(f"SELECT * FROM {entity['table']}")
    records = cursor.fetchall()

    cursor.close()
    db.close()

    return render_template(
        "entity.html",
        entity_key=entity_key,
        entity=entity,
        records=records
    )


@app.route("/entity/<entity_key>/add", methods=["POST"])
def add_entity(entity_key):
    if not is_logged_in():
        return redirect("/login")

    if entity_key not in ENTITIES:
        flash("Invalid entity.")
        return redirect("/dashboard")

    entity = ENTITIES[entity_key]
    columns = entity["columns"]

    values = [request.form[col] for col in columns]

    placeholders = ", ".join(["%s"] * len(columns))
    column_names = ", ".join(columns)

    db = get_db_connection()
    cursor = db.cursor()

    try:
        cursor.execute(
            f"INSERT INTO {entity['table']} ({column_names}) VALUES ({placeholders})",
            values
        )

        cursor.execute(
            "INSERT INTO activity_log (action) VALUES (%s)",
            (f"{entity['title']} record added",)
        )

        db.commit()
        flash("Record added successfully.")

    except Exception as e:
        db.rollback()
        flash("Error while adding record: " + str(e))

    cursor.close()
    db.close()

    return redirect(url_for("view_entity", entity_key=entity_key))


@app.route("/entity/<entity_key>/update/<pk_value>", methods=["POST"])
def update_entity(entity_key, pk_value):
    if not is_logged_in():
        return redirect("/login")

    if entity_key not in ENTITIES:
        flash("Invalid entity.")
        return redirect("/dashboard")

    entity = ENTITIES[entity_key]
    pk = entity["pk"]

    columns = [col for col in entity["columns"] if col != pk]

    set_clause = ", ".join([f"{col}=%s" for col in columns])
    values = [request.form[col] for col in columns]
    values.append(pk_value)

    db = get_db_connection()
    cursor = db.cursor()

    try:
        cursor.execute(
            f"UPDATE {entity['table']} SET {set_clause} WHERE {pk}=%s",
            values
        )

        cursor.execute(
            "INSERT INTO activity_log (action) VALUES (%s)",
            (f"{entity['title']} record updated",)
        )

        db.commit()
        flash("Record updated successfully.")

    except Exception as e:
        db.rollback()
        flash("Error while updating record: " + str(e))

    cursor.close()
    db.close()

    return redirect(url_for("view_entity", entity_key=entity_key))


@app.route("/entity/<entity_key>/delete/<pk_value>")
def delete_entity(entity_key, pk_value):
    if not is_logged_in():
        return redirect("/login")

    if entity_key not in ENTITIES:
        flash("Invalid entity.")
        return redirect("/dashboard")

    entity = ENTITIES[entity_key]
    pk = entity["pk"]

    db = get_db_connection()
    cursor = db.cursor()

    try:
        cursor.execute(
            f"DELETE FROM {entity['table']} WHERE {pk}=%s",
            (pk_value,)
        )

        cursor.execute(
            "INSERT INTO activity_log (action) VALUES (%s)",
            (f"{entity['title']} record deleted",)
        )

        db.commit()
        flash("Record deleted successfully.")

    except Exception as e:
        db.rollback()
        flash("Error while deleting record: " + str(e))

    cursor.close()
    db.close()

    return redirect(url_for("view_entity", entity_key=entity_key))


@app.route("/entity/<entity_key>/export")
def export_entity(entity_key):
    if not is_logged_in():
        return redirect("/login")

    if entity_key not in ENTITIES:
        flash("Invalid entity.")
        return redirect("/dashboard")

    entity = ENTITIES[entity_key]

    db = get_db_connection()
    cursor = db.cursor()

    cursor.execute(f"SELECT * FROM {entity['table']}")
    rows = cursor.fetchall()
    columns = [desc[0] for desc in cursor.description]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = entity["title"]

    sheet.append(columns)

    for row in rows:
        sheet.append(row)

    file_stream = BytesIO()
    workbook.save(file_stream)
    file_stream.seek(0)

    cursor.close()
    db.close()

    return send_file(
        file_stream,
        as_attachment=True,
        download_name=f"{entity_key}_records.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


if __name__ == "__main__":
    app.run(debug=True)



    

    db = get_db_connection()
cursor = db.cursor()

cursor.execute("SELECT * FROM shipment_booking")
records = cursor.fetchall()

cursor.close()
db.close()


